"""Convert the SRS markdown docs to Word (.docx).

Handles the constructs used in docs/srs/*.md: #/##/### headings,
GFM pipe tables, bullet/numbered lists, fenced code blocks, bold and
italic inline spans, and horizontal rules.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "word"
OUT.mkdir(exist_ok=True)


def add_runs(paragraph, text):
    """Add text to a paragraph, honoring **bold** and *italic* spans."""
    for token in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)


def style_base(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def add_table(doc, header, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(cell_text)
        run.bold = True
    for r, row in enumerate(rows, start=1):
        for c, cell_text in enumerate(row):
            cell = table.rows[r].cells[min(c, len(header) - 1)]
            cell.text = ""
            add_runs(cell.paragraphs[0], cell_text)
    doc.add_paragraph()


def convert(md_path: Path, out_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()

    doc = Document()
    style_base(doc)

    i = 0
    while i < len(lines):
        line = lines[i]

        # Fenced code block
        if line.strip().startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            continue

        # Table
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            add_table(doc, header, rows)
            continue

        # Horizontal rule
        if re.match(r"^\s*---+\s*$", line):
            doc.add_paragraph()
            i += 1
            continue

        # Headings
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=level)
            i += 1
            continue

        # Numbered list
        m = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(2))
            i += 1
            continue

        # Checkbox / bullet list
        m = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m:
            text = m.group(1)
            checkbox = re.match(r"^\[( |x)\]\s*(.*)$", text)
            if checkbox:
                mark = "☑" if checkbox.group(1) == "x" else "☐"
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(mark + "  ")
                add_runs(p, checkbox.group(2))
            else:
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, text)
            i += 1
            continue

        # Italic whole-line note
        if re.match(r"^\*[^*].*\*$", line.strip()):
            p = doc.add_paragraph()
            run = p.add_run(line.strip().strip("*"))
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        add_runs(p, line.strip())
        i += 1

    # Landscape-safe margins for wide tables
    for section in doc.sections:
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    sources = sorted((ROOT / "docs" / "srs").glob("*.md"))
    for src in sources:
        target = OUT / (src.stem + ".docx")
        convert(src, target)
        print(f"converted {src.name} -> {target.relative_to(ROOT)}")
